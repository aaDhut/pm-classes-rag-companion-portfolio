"""Turn each downloaded file into a list of units: {"page": int | None, "text": str, "via": str}.

PDF pages (slide decks, and Google Docs exported as PDF):
  1. text layer from PyMuPDF
  2. Google Cloud Vision OCR - the whole page if it is an image slide, otherwise
     each large embedded image (screenshots carry text the text layer misses)
  3. Gemini "Visual:" description for pages with diagrams/screenshots

Results are cached in data/extracted/<file id>.json keyed by Drive modifiedTime and
EXTRACT_VERSION, so OCR and Gemini run once per file version.
"""
import csv
import io
import json
import posixpath
import re
import zipfile
from xml.etree import ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pymupdf

from . import config, llm, ocr

IMAGE_EXT = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
MIN_IMAGE_SHARE = 0.05   # ignore icons/logos smaller than this share of the page
MIN_DRAWINGS = 25        # vector diagrams (boxes/arrows drawn in the slide) also count as visual
OCR_DPI, VISUAL_DPI = 150, 110


def _render(page, dpi, clip=None) -> bytes:
    return page.get_pixmap(dpi=dpi, clip=clip).tobytes("jpeg", jpg_quality=85)


def _page_plan(page) -> dict:
    area = page.rect.width * page.rect.height
    boxes = []
    for info in page.get_image_info():
        r = pymupdf.Rect(info["bbox"]) & page.rect
        if not r.is_empty and r.width * r.height / area >= MIN_IMAGE_SHARE:
            boxes.append(r)
    image_share = min(sum(r.width * r.height for r in boxes) / area, 1.0)
    text = page.get_text("text").strip()
    whole_page_ocr = len(text) < config.VISION_MIN_CHARS
    return {
        "text": text,
        "ocr_boxes": [] if whole_page_ocr else boxes,
        "whole_page_ocr": whole_page_ocr,
        "visual": whole_page_ocr or image_share >= config.VISUAL_MIN_AREA
                  or len(page.get_drawings()) >= MIN_DRAWINGS,
    }


def _pdf(path: Path, label: str) -> list[dict]:
    doc = pymupdf.open(path)
    plans = [_page_plan(p) for p in doc]

    # 1) Cloud Vision OCR, batched across the whole file
    jobs = []  # (page index, image bytes)
    for i, plan in enumerate(plans):
        if plan["whole_page_ocr"]:
            jobs.append((i, _render(doc[i], OCR_DPI)))
        else:
            jobs += [(i, _render(doc[i], OCR_DPI, clip=r)) for r in plan["ocr_boxes"]]
    image_text = {i: [] for i in range(len(plans))}
    if jobs:
        print(f"    Cloud Vision OCR: {len(jobs)} image(s) in {label}")
        for (i, _), text in zip(jobs, ocr.ocr_images([img for _, img in jobs])):
            if len(text) >= 15:
                image_text[i].append(text)

    units = []
    for i, plan in enumerate(plans):
        if plan["whole_page_ocr"]:
            ocr_text = "\n".join(image_text[i])
            text, via = (ocr_text, "ocr") if len(ocr_text) > len(plan["text"]) else (plan["text"], "text")
        else:
            text, via = plan["text"], "text"
            if image_text[i]:
                text += "\n\n[Text in screenshots/images]\n" + "\n---\n".join(image_text[i])
                via = "text+ocr"
        units.append({"page": i + 1, "text": text.strip(), "via": via})

    # 2) Gemini explains diagrams/screenshots
    visual_pages = [i for i, p in enumerate(plans) if p["visual"]]
    if visual_pages:
        print(f"    Gemini visual descriptions: {len(visual_pages)} page(s)")

        def describe(i):
            try:
                return i, llm.describe_visual(_render(doc[i], VISUAL_DPI), units[i]["text"], "image/jpeg")
            except Exception as e:
                print(f"    ! Gemini failed on page {i + 1}: {e}")
                return i, ""

        with ThreadPoolExecutor(6) as pool:
            for i, visual in pool.map(describe, visual_pages):
                if visual:
                    units[i]["text"] = f"{units[i]['text']}\n\nVisual:\n{visual}".strip()
                    units[i]["via"] += "+visual"
    return [u for u in units if u["text"]]


def _image(path: Path, ext: str) -> list[dict]:
    data = path.read_bytes()
    text = ocr.ocr_images([data])[0]
    visual = llm.describe_visual(data, text, IMAGE_EXT[ext])
    body = "\n\n".join(x for x in [text, f"Visual:\n{visual}" if visual else ""] if x)
    return [{"page": None, "text": body, "via": "ocr+visual"}] if body else []


def _pictures(pictures: list[tuple[bytes, str]], context: str) -> str:
    """OCR + Gemini description for pictures embedded in Word/Excel files, as one text block."""
    if not pictures:
        return ""
    texts = ocr.ocr_images([b for b, _ in pictures])

    def describe(k):
        blob, mime = pictures[k]
        try:
            return llm.describe_visual(blob, f"{texts[k]}\n\n{context}"[:4000], mime)
        except Exception as e:  # e.g. EMF/WMF images Gemini can't read
            print(f"    ! Gemini failed on picture {k + 1}: {e}")
            return ""

    with ThreadPoolExecutor(6) as pool:
        visuals = list(pool.map(describe, range(len(pictures))))
    blocks = []
    for k, (text, visual) in enumerate(zip(texts, visuals), 1):
        body = "\n".join(x for x in [text if len(text) >= 15 else "", f"Visual:\n{visual}" if visual else ""] if x)
        if body:
            blocks.append(f"[Picture {k}]\n{body}")
    return "\n\n".join(blocks)


def _mime(name: str) -> str:
    return IMAGE_EXT.get(Path(name).suffix.lower(), "image/" + Path(name).suffix.lower().lstrip("."))


def _docx(path: Path) -> list[dict]:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    pictures = [(rel.target_part.blob, rel.target_part.content_type)
                for rel in d.part.rels.values() if "image" in rel.reltype]
    pictures = [p for p in pictures if len(p[0]) > 30_000]  # skip icons
    via = "text"
    if pictures:
        print(f"    Pictures (OCR + Gemini): {len(pictures)} in {path.name}")
        if block := _pictures(pictures, "\n".join(parts)[:2000]):
            parts.append(block)
            via = "text+ocr+visual"
    return [{"page": None, "text": "\n".join(parts), "via": via}]


def _pptx(path: Path) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    units = []
    for i, slide in enumerate(Presentation(path).slides):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        pictures = [sh.image for sh in slide.shapes
                    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and len(sh.image.blob) > 30_000]
        via = "text"
        if pictures:
            found = [t for t in ocr.ocr_images([p.blob for p in pictures]) if len(t) >= 15]
            if found:
                texts.append("[Text in screenshots/images]\n" + "\n---\n".join(found))
                via = "text+ocr"
            biggest = max(pictures, key=lambda p: len(p.blob))
            try:
                visual = llm.describe_visual(biggest.blob, "\n".join(texts), biggest.content_type)
            except Exception as e:  # e.g. EMF/WMF images Gemini can't read
                print(f"    ! Gemini failed on slide {i + 1}: {e}")
                visual = ""
            if visual:
                texts.append(f"Visual:\n{visual}")
                via += "+visual"
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        if texts:
            units.append({"page": i + 1, "text": "\n".join(texts), "via": via})
    return units


C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NS = {"c": C_NS, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
MIN_SHEET_PICTURE = 5_000  # bytes; smaller is a logo/icon


def _rels(z: zipfile.ZipFile, part: str) -> list[tuple[str, str]]:
    """(relationship type, target zip path) for one part of an Office file."""
    rels = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rels not in z.namelist():
        return []
    out = []
    for r in ET.fromstring(z.read(rels)):
        target = r.get("Target", "")
        target = target.lstrip("/") if target.startswith("/") else \
            posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
        out.append((r.get("Type", ""), target))
    return out


def _sheet_drawings(z: zipfile.ZipFile) -> dict[str, list[str]]:
    """Sheet title -> chart and picture parts drawn on that sheet."""
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    targets = {}
    for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")):
        t = r.get("Target", "")
        targets[r.get("Id")] = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join("xl", t))
    out = {}
    for sh in wb.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
        rid = sh.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        parts = []
        for kind, drawing in _rels(z, targets.get(rid, "")):
            if kind.endswith("/drawing"):
                parts += [t for k, t in _rels(z, drawing) if k.endswith("/chart") or k.endswith("/image")]
        out[sh.get("name")] = parts
    return out


def _chart_text(xml: bytes) -> str:
    """Title, type, axis titles and every series' values, read from the chart's cached data."""
    root = ET.fromstring(xml)
    title = "".join(t.text or "" for t in root.findall("c:chart/c:title//a:t", NS)).strip()
    plot = root.find("c:chart/c:plotArea", NS)
    if plot is None:
        return ""
    kinds = [re.sub(r"(?<!^)([A-Z])", r" \1", el.tag.split("}")[1]).lower() for el in plot if el.tag.endswith("Chart")]
    lines = [f"Chart: {title or '(untitled)'} ({', '.join(kinds) or 'chart'})"]
    for ax in plot:
        if ax.tag.endswith("Ax"):
            ax_title = "".join(t.text or "" for t in ax.findall("c:title//a:t", NS)).strip()
            if ax_title:
                lines.append(f"Axis: {ax_title}")

    def points(el):
        return {int(pt.get("idx")): pt.findtext("c:v", "", NS) for pt in el.iter(f"{{{C_NS}}}pt")} if el is not None else {}

    for ser in plot.iter(f"{{{C_NS}}}ser"):
        name = " ".join(points(ser.find("c:tx", NS)).values()) or "series"
        cats = points(ser.find("c:cat", NS)) or points(ser.find("c:xVal", NS))
        vals = points(ser.find("c:val", NS)) or points(ser.find("c:yVal", NS))
        pairs = ", ".join(f"{cats[i]}: {v}" if i in cats else v for i, v in sorted(vals.items()))
        lines.append(f"Series '{name}': {pairs}")
    return "\n".join(lines)


def _xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    data = path.read_bytes()
    # Opened as a stream: openpyxl rejects paths without an .xlsx suffix, and uploaded Drive files have none.
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    z = zipfile.ZipFile(io.BytesIO(data))
    drawings = _sheet_drawings(z)
    units = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(" | ".join(cells))
            if len(rows) >= 400:
                break
        text, via = (f"Sheet: {ws.title}\n" + "\n".join(rows)) if rows else f"Sheet: {ws.title}", "text"
        parts = drawings.get(ws.title, [])
        charts = [c for c in (_chart_text(z.read(p)) for p in parts if "/charts/" in p) if c]
        if charts:
            print(f"    Charts read: {len(charts)} on sheet '{ws.title}'")
            text += "\n\n[Charts]\n" + "\n\n".join(charts)
            via += "+chart"
        pictures = [(z.read(p), _mime(p)) for p in parts if "/media/" in p and z.getinfo(p).file_size >= MIN_SHEET_PICTURE]
        if pictures:
            print(f"    Pictures (OCR + Gemini): {len(pictures)} on sheet '{ws.title}'")
            if block := _pictures(pictures, text[:2000]):
                text += "\n\n" + block
                via += "+ocr+visual"
        if rows or via != "text":
            units.append({"page": None, "text": text, "via": via})
    return units


def _csv(path: Path) -> list[dict]:
    # Datasets (e.g. event logs) can be huge; the header and a sample are what matter for Q&A.
    with open(path, newline="", encoding="utf-8", errors="replace") as fh:
        rows = [" | ".join(r) for _, r in zip(range(41), csv.reader(fh))]
    total = sum(1 for _ in open(path, encoding="utf-8", errors="replace")) - 1
    text = f"Dataset with {total} rows. Columns: {rows[0] if rows else ''}\nSample rows:\n" + "\n".join(rows[1:])
    return [{"page": None, "text": text, "via": "text"}]


def _text(path: Path) -> list[dict]:
    return [{"page": None, "text": path.read_text(encoding="utf-8", errors="replace").replace("﻿", ""), "via": "text"}]


def _sniff_ext(path: Path) -> str:
    """Drive files sometimes lack an extension; guess from the first bytes."""
    head = path.read_bytes()[:8]
    if head.startswith(b"%PDF"):
        return ".pdf"
    if head.startswith(b"\x89PNG"):
        return ".png"
    if head.startswith(b"\xff\xd8"):
        return ".jpg"
    return path.suffix.lower()


def extract(entry: dict) -> list[dict]:
    """The file's units; PDF pages also get their tables as rows and columns (rag/tables.py)."""
    units = _units(entry)
    path = config.ROOT / entry["path"]
    if units and _sniff_ext(path) == ".pdf":
        from . import tables
        units = tables.add_to_units(entry, path, units)
    return units


def _units(entry: dict) -> list[dict]:
    cache = config.EXTRACT_DIR / f"{entry['id']}.json"
    if cache.exists():
        cached = json.loads(cache.read_text())
        if cached["modifiedTime"] == entry["modifiedTime"] and cached.get("version") == config.EXTRACT_VERSION \
                and cached.get("path") == entry["path"]:
            return cached["units"]
    path = config.ROOT / entry["path"]
    ext = _sniff_ext(path)
    name = entry["name"].lower()
    if ext == ".pdf":
        units = _pdf(path, entry["name"])
    elif ext == ".docx" or name.endswith(".docx"):
        units = _docx(path)
    elif ext == ".pptx" or name.endswith(".pptx"):
        units = _pptx(path)
    elif ext == ".xlsx" or name.endswith(".xlsx"):
        units = _xlsx(path)
    elif ext == ".csv" or name.endswith(".csv"):
        units = _csv(path)
    elif ext in IMAGE_EXT:
        units = _image(path, ext)
    elif ext in (".txt", ".md") or entry["mimeType"].startswith("text/"):
        units = _text(path)
    else:
        print(f"    - no extractor for {entry['name']} ({entry['mimeType']})")
        units = []
    config.EXTRACT_DIR.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps({"modifiedTime": entry["modifiedTime"], "version": config.EXTRACT_VERSION,
                                 "path": entry["path"], "units": units}))
    return units
